"""Reading history export — CSV, Excel, PDF report, and JSON backup.

Routes validate query params and stream the returned buffer. All filtering,
formatting, and file construction lives here. Read-only: never writes to the DB.
"""
from __future__ import annotations

import csv
import io
import json
import os
from collections import Counter, defaultdict
from datetime import date, datetime
from typing import Any

from domain.reading_status import ReadingStatus
from utils.ratings import has_rating

# Optional heavy deps — imported lazily in export methods so unit-style
# imports of this module stay light when packages are missing in tests.


BACKUP_FORMAT = 'kwalitec-library-backup'
BACKUP_VERSION = 1

CSV_HEADERS = [
    'Book Title',
    'Author',
    'Rating',
    'Reading Status',
    'Wild Book',
    'Date Added',
    'Started Reading',
    'Finished Reading',
    'Reading Duration',
    'Quote Count',
    'Favourite Quote',
    'Cover Present',
    'Finished Year',
    'Finished Month',
]

STATUS_FILTER_ALL = 'all'
STATUS_FILTER_WILD = 'wild'

VALID_EXPORT_TYPES = frozenset({'csv', 'excel', 'pdf', 'json'})
VALID_STATUS_FILTERS = frozenset({
    STATUS_FILTER_ALL,
    ReadingStatus.FINISHED,
    ReadingStatus.READING,
    ReadingStatus.WANT_TO_READ,
    STATUS_FILTER_WILD,
})


class ExportError(Exception):
    """User-safe export failure."""

    def __init__(self, message: str):
        self.message = message
        super().__init__(message)


class ExportService:
    """Build downloadable reading-history archives for one user."""

    def __init__(self, reading_service: Any, statistics_service: Any):
        self.reading_service = reading_service
        self.statistics_service = statistics_service

    # ── Public API ────────────────────────────────────────────

    def export(
        self,
        *,
        export_type: str,
        books: list[Any],
        goals: list[Any],
        username: str,
        user_created_at: datetime | None = None,
        avatar: str | None = None,
        status: str | None = None,
        author: str | None = None,
        rating: int | None = None,
        year: int | None = None,
        theme: str = 'light',
        now: datetime | None = None,
    ) -> tuple[io.BytesIO, str, str]:
        """Return (buffer, download_filename, mimetype).

        JSON backup always uses the full library (filters ignored).
        Other formats apply optional filters to the book set.
        """
        now = now or datetime.utcnow()
        export_type = (export_type or '').strip().lower()
        if export_type not in VALID_EXPORT_TYPES:
            raise ExportError('Please choose a valid export type.')

        if not books:
            raise ExportError('No books available to export.')

        theme = 'dark' if (theme or '').strip().lower() == 'dark' else 'light'
        stamp = now.strftime('%Y-%m-%d')

        if export_type == 'json':
            return self._export_json_backup(
                books=books,
                goals=goals,
                username=username,
                user_created_at=user_created_at,
                avatar=avatar,
                now=now,
                stamp=stamp,
            )

        filtered = self.apply_filters(
            books,
            status=status,
            author=author,
            rating=rating,
            year=year,
        )
        if not filtered:
            raise ExportError('No books match the selected filters.')

        if export_type == 'csv':
            return self._export_csv(filtered, stamp)
        if export_type == 'excel':
            return self._export_excel(
                filtered,
                books_all=books,
                goals=goals,
                username=username,
                now=now,
                stamp=stamp,
            )
        return self._export_pdf(
            filtered,
            books_all=books,
            goals=goals,
            username=username,
            now=now,
            stamp=stamp,
            theme=theme,
        )

    def apply_filters(
        self,
        books: list[Any],
        *,
        status: str | None = None,
        author: str | None = None,
        rating: int | None = None,
        year: int | None = None,
    ) -> list[Any]:
        """Filter books in memory (caller already scoped to current user)."""
        status_key = (status or STATUS_FILTER_ALL).strip().lower()
        if status_key not in VALID_STATUS_FILTERS:
            status_key = STATUS_FILTER_ALL

        author_key = (author or '').strip()
        result = []
        for book in books:
            if status_key == STATUS_FILTER_WILD:
                if not bool(getattr(book, 'is_wild', False)):
                    continue
            elif status_key != STATUS_FILTER_ALL:
                book_status = (
                    book.reading_status
                    if ReadingStatus.is_valid(book.reading_status)
                    else ReadingStatus.FINISHED
                )
                if book_status != status_key:
                    continue

            if author_key and book.author != author_key:
                continue

            if rating is not None:
                if not has_rating(book) or int(book.rating) != int(rating):
                    continue

            if year is not None:
                finished = getattr(book, 'finished_reading_at', None)
                if finished is None or finished.year != int(year):
                    continue

            result.append(book)
        return result

    def filter_options(self, books: list[Any]) -> dict[str, Any]:
        """Authors and finished years for the Settings export form."""
        authors = sorted({b.author for b in books if b.author}, key=str.casefold)
        years: set[int] = set()
        for b in books:
            finished = getattr(b, 'finished_reading_at', None)
            if finished is not None:
                years.add(finished.year)
        return {
            'authors': authors,
            'years': sorted(years, reverse=True),
            'book_count': len(books),
        }

    # ── Row builders ──────────────────────────────────────────

    def build_history_rows(self, books: list[Any]) -> list[dict[str, Any]]:
        """One display/export row per book (quotes already selectinloaded)."""
        rows = []
        for book in books:
            quotes = list(getattr(book, 'quotes', None) or [])
            favourite = next((q for q in quotes if getattr(q, 'is_favourite', False)), None)
            duration = self.reading_service.calculate_duration(
                getattr(book, 'started_reading_at', None),
                getattr(book, 'finished_reading_at', None),
            )
            finished = getattr(book, 'finished_reading_at', None)
            status = (
                book.reading_status
                if ReadingStatus.is_valid(book.reading_status)
                else ReadingStatus.FINISHED
            )
            cover_present = bool(
                getattr(book, 'cover_data', None) or getattr(book, 'cover_filename', None)
            )
            rows.append({
                'title': book.title or '',
                'author': book.author or '',
                'rating': int(book.rating) if has_rating(book) else '',
                'reading_status': ReadingStatus.label(status),
                'reading_status_key': status,
                'is_wild': 'Yes' if getattr(book, 'is_wild', False) else 'No',
                'is_wild_bool': bool(getattr(book, 'is_wild', False)),
                'date_added': self._fmt_datetime(getattr(book, 'date_added', None)),
                'started_reading': self._fmt_date(getattr(book, 'started_reading_at', None)),
                'finished_reading': self._fmt_date(finished),
                'reading_duration': (
                    f'{duration} day{"s" if duration != 1 else ""}'
                    if duration is not None else ''
                ),
                'reading_duration_days': duration,
                'quote_count': len(quotes),
                'favourite_quote': (favourite.text if favourite else '') or '',
                'cover_present': 'Yes' if cover_present else 'No',
                'finished_year': finished.year if finished else '',
                'finished_month': finished.strftime('%B') if finished else '',
                'finished_month_num': finished.month if finished else None,
                'book': book,
            })
        return rows

    def build_summary(
        self,
        books: list[Any],
        *,
        goals: list[Any],
        username: str,
        now: datetime,
        scope_label: str = 'Filtered export',
    ) -> dict[str, Any]:
        """Reading summary metrics for Excel sheet 2 / PDF."""
        dash = self.statistics_service.reading_dashboard(books, self.reading_service, now=now)
        kpis = dash['kpis']
        today = now.date() if isinstance(now, datetime) else now
        year = today.year
        finished_this_year = int(
            (dash.get('reading_summary') or {}).get('finished_this_year') or 0
        )
        goal = next((g for g in goals if getattr(g, 'year', None) == year), None)
        goal_progress = self.statistics_service.annual_goal_progress(
            goal,
            finished_this_year,
            today=today,
        )
        goal_display = (
            f"{goal_progress['completed']} / {goal_progress['target']}"
            if goal_progress.get('exists')
            else 'Not set'
        )
        return {
            'username': username,
            'export_date': now.strftime('%Y-%m-%d'),
            'scope_label': scope_label,
            'books': len(books),
            'finished': kpis.get('books_finished', 0),
            'currently_reading': kpis.get('currently_reading', 0),
            'want_to_read': kpis.get('want_to_read', 0),
            'average_rating': kpis.get('average_rating_display') or '—',
            'average_rating_value': kpis.get('average_rating'),
            'wild_books': kpis.get('wild_books', 0),
            'favourite_author': kpis.get('favourite_author') or '—',
            'current_goal': goal_display,
            'goal_target': goal_progress.get('target'),
            'goal_completed': goal_progress.get('completed', 0),
            'finished_this_year': finished_this_year,
            'rating_distribution': dash['charts']['rating_distribution'],
            'finished_per_month': dash['charts']['finished_per_month'],
            'activity': self.reading_service.recent_activity(books, limit=20),
        }

    # ── CSV ───────────────────────────────────────────────────

    def _export_csv(self, books: list[Any], stamp: str) -> tuple[io.BytesIO, str, str]:
        rows = self.build_history_rows(books)
        text = io.StringIO(newline='')
        writer = csv.writer(text, dialect='excel', quoting=csv.QUOTE_MINIMAL, lineterminator='\r\n')
        writer.writerow(CSV_HEADERS)
        for row in rows:
            writer.writerow([
                row['title'],
                row['author'],
                row['rating'],
                row['reading_status'],
                row['is_wild'],
                row['date_added'],
                row['started_reading'],
                row['finished_reading'],
                row['reading_duration'],
                row['quote_count'],
                row['favourite_quote'],
                row['cover_present'],
                row['finished_year'],
                row['finished_month'],
            ])
        # UTF-8 with BOM for Excel compatibility on Windows
        payload = ('\ufeff' + text.getvalue()).encode('utf-8')
        buf = io.BytesIO(payload)
        buf.seek(0)
        filename = f'kwalitec-library-history-{stamp}.csv'
        return buf, filename, 'text/csv; charset=utf-8'

    # ── Excel ─────────────────────────────────────────────────

    def _export_excel(
        self,
        books: list[Any],
        *,
        books_all: list[Any],
        goals: list[Any],
        username: str,
        now: datetime,
        stamp: str,
    ) -> tuple[io.BytesIO, str, str]:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        rows = self.build_history_rows(books)
        summary = self.build_summary(
            books,
            goals=goals,
            username=username,
            now=now,
            scope_label='Export selection',
        )

        wb = Workbook()
        ws = wb.active
        ws.title = 'Reading History'

        header_fill = PatternFill('solid', fgColor='5B54E8')
        header_font = Font(bold=True, color='FFFFFF', name='Calibri', size=11)
        cell_font = Font(name='Calibri', size=11)
        thin = Border(
            left=Side(style='thin', color='E5E7EB'),
            right=Side(style='thin', color='E5E7EB'),
            top=Side(style='thin', color='E5E7EB'),
            bottom=Side(style='thin', color='E5E7EB'),
        )
        alt_fill = PatternFill('solid', fgColor='F5F3FF')

        for col, header in enumerate(CSV_HEADERS, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        for r_idx, row in enumerate(rows, 2):
            values = [
                row['title'],
                row['author'],
                row['rating'] if row['rating'] != '' else None,
                row['reading_status'],
                row['is_wild'],
                row['date_added'] or None,
                row['started_reading'] or None,
                row['finished_reading'] or None,
                row['reading_duration'] or None,
                row['quote_count'],
                row['favourite_quote'] or None,
                row['cover_present'],
                row['finished_year'] if row['finished_year'] != '' else None,
                row['finished_month'] or None,
            ]
            for c_idx, value in enumerate(values, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                cell.font = cell_font
                cell.border = thin
                cell.alignment = Alignment(vertical='center', wrap_text=c_idx in (1, 2, 11))
                if r_idx % 2 == 0:
                    cell.fill = alt_fill

        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{get_column_letter(len(CSV_HEADERS))}{max(1, len(rows) + 1)}'
        self._autosize_columns(ws, max_width=48)

        # Sheet 2 — Reading Summary
        ws2 = wb.create_sheet('Reading Summary')
        title_font = Font(bold=True, name='Calibri', size=16, color='5B54E8')
        label_font = Font(bold=True, name='Calibri', size=11, color='374151')
        value_font = Font(name='Calibri', size=11)

        ws2['A1'] = 'Kwalitec Library — Reading Summary'
        ws2['A1'].font = title_font
        ws2.merge_cells('A1:B1')

        summary_rows = [
            ('Username', summary['username']),
            ('Export Date', summary['export_date']),
            ('Books', summary['books']),
            ('Finished', summary['finished']),
            ('Currently Reading', summary['currently_reading']),
            ('Want To Read', summary['want_to_read']),
            ('Average Rating', summary['average_rating']),
            ('Wild Books', summary['wild_books']),
            ('Favourite Author', summary['favourite_author']),
            ('Current Goal', summary['current_goal']),
            ('Finished This Year', summary['finished_this_year']),
            ('Library Size (all)', len(books_all)),
        ]
        for i, (label, value) in enumerate(summary_rows, 3):
            a = ws2.cell(row=i, column=1, value=label)
            b = ws2.cell(row=i, column=2, value=value)
            a.font = label_font
            b.font = value_font
            a.border = thin
            b.border = thin

        ws2.column_dimensions['A'].width = 22
        ws2.column_dimensions['B'].width = 36

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f'kwalitec-library-history-{stamp}.xlsx'
        return (
            buf,
            filename,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    @staticmethod
    def _autosize_columns(ws, *, max_width: int = 48) -> None:
        from openpyxl.utils import get_column_letter

        for col_cells in ws.columns:
            letter = get_column_letter(col_cells[0].column)
            longest = 0
            for cell in col_cells:
                if cell.value is None:
                    continue
                longest = max(longest, min(len(str(cell.value)), max_width))
            ws.column_dimensions[letter].width = min(max(longest + 2, 10), max_width)

    # ── PDF (Pillow multi-page; matches ShareEngine branding) ──

    _PDF_W = 850
    _PDF_H = 1202
    _PDF_MARGIN = 48

    _FONT_REG_PATHS = [
        '/System/Library/Fonts/Supplemental/Arial.ttf',
        '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
        '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf',
        '/System/Library/Fonts/Supplemental/Arial Unicode.ttf',
        'C:/Windows/Fonts/arial.ttf',
    ]
    _FONT_BOLD_PATHS = [
        '/System/Library/Fonts/Supplemental/Arial Bold.ttf',
        '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',
        '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf',
        'C:/Windows/Fonts/arialbd.ttf',
    ]

    def _export_pdf(
        self,
        books: list[Any],
        *,
        books_all: list[Any],
        goals: list[Any],
        username: str,
        now: datetime,
        stamp: str,
        theme: str,
    ) -> tuple[io.BytesIO, str, str]:
        from PIL import Image, ImageDraw, ImageFont

        palette = self._pdf_palette(theme)
        fonts = self._pdf_fonts(ImageFont)
        summary = self.build_summary(
            books,
            goals=goals,
            username=username,
            now=now,
            scope_label='Reading report',
        )
        rows = self.build_history_rows(books)
        top_rated = sorted(
            [r for r in rows if r['rating'] != ''],
            key=lambda r: (-int(r['rating']), r['title'].casefold()),
        )[:8]
        favourite_quote = next(
            (r['favourite_quote'] for r in rows if r['favourite_quote']),
            '',
        )
        recent = sorted(
            books,
            key=lambda b: getattr(b, 'date_added', None) or datetime.min,
            reverse=True,
        )[:8]
        month_labels, month_values = self._month_chart_series(books, now)
        year_labels, year_values = self._year_chart_series(books)
        rating_values = self._chart_values(summary['rating_distribution'])

        pages: list[Any] = []
        pages.append(self._pdf_page_cover(
            Image, ImageDraw, fonts, palette,
            username=username, export_date=summary['export_date'],
        ))
        pages.append(self._pdf_page_summary(
            Image, ImageDraw, fonts, palette, summary, len(books_all),
        ))
        pages.append(self._pdf_page_top_rated(
            Image, ImageDraw, fonts, palette, top_rated,
        ))
        pages.append(self._pdf_page_quote(
            Image, ImageDraw, fonts, palette, favourite_quote,
        ))
        pages.append(self._pdf_page_timeline(
            Image, ImageDraw, fonts, palette, summary['activity'][:12],
        ))
        pages.append(self._pdf_page_recent(
            Image, ImageDraw, fonts, palette, recent,
        ))
        pages.append(self._pdf_page_charts(
            Image, ImageDraw, fonts, palette,
            rating_values=rating_values,
            month_labels=month_labels,
            month_values=month_values,
            year_labels=year_labels,
            year_values=year_values,
        ))

        buf = io.BytesIO()
        first, rest = pages[0], pages[1:]
        first.save(
            buf,
            format='PDF',
            save_all=True,
            append_images=rest,
            resolution=100.0,
            quality=85,
        )
        buf.seek(0)
        return buf, f'kwalitec-library-report-{stamp}.pdf', 'application/pdf'

    def _pdf_fonts(self, ImageFont: Any) -> dict[str, Any]:
        reg_path = next((p for p in self._FONT_REG_PATHS if os.path.isfile(p)), None)
        bold_path = next((p for p in self._FONT_BOLD_PATHS if os.path.isfile(p)), None)

        def load(size: int, bold: bool = False):
            path = bold_path if bold and bold_path else reg_path
            if path:
                try:
                    return ImageFont.truetype(path, size)
                except OSError:
                    pass
            return ImageFont.load_default()

        return {
            'title': load(28, True),
            'h1': load(22, True),
            'h2': load(18, True),
            'body': load(14),
            'body_bold': load(14, True),
            'small': load(12),
            'tiny': load(11),
            'metric': load(18, True),
            'brand': load(26, True),
        }

    @staticmethod
    def _pdf_palette(theme: str) -> dict[str, tuple[int, int, int]]:
        if theme == 'dark':
            # Midnight OLED — match design-system dark tokens
            return {
                'bg': (0, 0, 0),
                'card': (18, 18, 20),
                'text': (184, 184, 188),
                'muted': (128, 128, 134),
                'primary': (101, 94, 245),
                'primary_dark': (79, 72, 212),
                'accent': (246, 200, 95),
                'border': (31, 31, 34),
                'on_primary': (255, 255, 255),
            }
        # Light E-Ink — match design-system light tokens
        return {
            'bg': (237, 238, 239),
            'card': (250, 250, 249),
            'text': (10, 10, 10),
            'muted': (82, 82, 91),
            'primary': (91, 84, 232),
            'primary_dark': (67, 56, 202),
            'accent': (246, 200, 95),
            'border': (212, 212, 216),
            'on_primary': (255, 255, 255),
        }

    def _pdf_blank(self, Image: Any, palette: dict) -> Any:
        return Image.new('RGB', (self._PDF_W, self._PDF_H), palette['bg'])

    def _pdf_draw_section_chrome(
        self, draw: Any, fonts: dict, palette: dict, title: str
    ) -> int:
        m = self._PDF_MARGIN
        draw.text((m, 56), title, fill=palette['primary'], font=fonts['h1'])
        draw.rectangle((m, 100, m + 140, 104), fill=palette['primary'])
        brand = 'Kwalitec Library'
        bw = draw.textlength(brand, font=fonts['tiny'])
        draw.text((self._PDF_W - m - bw, 64), brand, fill=palette['muted'], font=fonts['tiny'])
        return 140

    def _pdf_page_cover(
        self, Image, ImageDraw, fonts, palette, *, username, export_date
    ):
        img = self._pdf_blank(Image, palette)
        draw = ImageDraw.Draw(img)
        w, h = self._PDF_W, self._PDF_H
        draw.rectangle((0, 0, w, 180), fill=palette['primary'])
        draw.rectangle((0, h - 70, w, h), fill=palette['primary_dark'])

        brand = 'Kwalitec Library'
        bw = draw.textlength(brand, font=fonts['brand'])
        draw.text(((w - bw) / 2, 70), brand, fill=palette['on_primary'], font=fonts['brand'])
        sub = 'Reading Journal'
        sw = draw.textlength(sub, font=fonts['body'])
        draw.text(((w - sw) / 2, 110), sub, fill=palette['on_primary'], font=fonts['body'])

        title = 'Reading Report'
        tw = draw.textlength(title, font=fonts['title'])
        draw.text(((w - tw) / 2, h / 2 - 30), title, fill=palette['text'], font=fonts['title'])
        uw = draw.textlength(username, font=fonts['h2'])
        draw.text(((w - uw) / 2, h / 2 + 20), username, fill=palette['muted'], font=fonts['h2'])
        date_line = f'Exported {export_date}'
        dw = draw.textlength(date_line, font=fonts['body'])
        draw.text(((w - dw) / 2, h / 2 + 52), date_line, fill=palette['muted'], font=fonts['body'])

        foot = 'Your reading history · Yours forever'
        fw = draw.textlength(foot, font=fonts['small'])
        draw.text(((w - fw) / 2, h - 42), foot, fill=palette['on_primary'], font=fonts['small'])
        return img

    def _pdf_page_summary(self, Image, ImageDraw, fonts, palette, summary, library_total):
        img = self._pdf_blank(Image, palette)
        draw = ImageDraw.Draw(img)
        y = self._pdf_draw_section_chrome(draw, fonts, palette, 'Reading Summary')
        metrics = [
            ('Books', str(summary['books'])),
            ('Average Rating', str(summary['average_rating'])),
            ('Favourite Author', str(summary['favourite_author'])),
            ('Wild Books', str(summary['wild_books'])),
            ('Reading Goal', str(summary['current_goal'])),
            ('Finished This Year', str(summary['finished_this_year'])),
        ]
        m = self._PDF_MARGIN
        col_w = (self._PDF_W - 2 * m - 16) // 2
        card_h = 78
        for i, (label, value) in enumerate(metrics):
            col = i % 2
            row = i // 2
            x = m + col * (col_w + 16)
            yy = y + row * (card_h + 14)
            draw.rounded_rectangle((x, yy, x + col_w, yy + card_h), radius=12, fill=palette['card'])
            draw.text((x + 16, yy + 14), label.upper(), fill=palette['muted'], font=fonts['tiny'])
            draw.text(
                (x + 16, yy + 38),
                self._truncate(str(value), 28),
                fill=palette['text'],
                font=fonts['metric'],
            )
        note_y = y + 3 * (card_h + 14) + 16
        draw.text(
            (m, note_y),
            f'Library total: {library_total} books · Export scope shown above',
            fill=palette['muted'],
            font=fonts['small'],
        )
        return img

    def _pdf_page_top_rated(self, Image, ImageDraw, fonts, palette, top_rated):
        img = self._pdf_blank(Image, palette)
        draw = ImageDraw.Draw(img)
        y = self._pdf_draw_section_chrome(draw, fonts, palette, 'Top Rated Books')
        m = self._PDF_MARGIN
        if not top_rated:
            draw.text((m, y), 'No rated books in this export.', fill=palette['muted'], font=fonts['body'])
            return img
        for i, row in enumerate(top_rated, 1):
            stars = f"{int(row['rating'])}/5"
            draw.rounded_rectangle(
                (m, y, self._PDF_W - m, y + 64), radius=12, fill=palette['card']
            )
            draw.text((m + 18, y + 14), f'{i}. {self._truncate(row["title"], 42)}',
                      fill=palette['text'], font=fonts['body_bold'])
            draw.text((m + 18, y + 38), row['author'], fill=palette['muted'], font=fonts['small'])
            sw = draw.textlength(stars, font=fonts['body_bold'])
            draw.text((self._PDF_W - m - 18 - sw, y + 24), stars,
                      fill=palette['accent'], font=fonts['body_bold'])
            y += 76
        return img

    def _pdf_page_quote(self, Image, ImageDraw, fonts, palette, favourite_quote):
        img = self._pdf_blank(Image, palette)
        draw = ImageDraw.Draw(img)
        y = self._pdf_draw_section_chrome(draw, fonts, palette, 'Favourite Quote')
        m = self._PDF_MARGIN
        box_h = 260
        draw.rounded_rectangle(
            (m, y, self._PDF_W - m, y + box_h), radius=14, fill=palette['card']
        )
        draw.rectangle((m, y, m + 8, y + box_h), fill=palette['primary'])
        if favourite_quote:
            lines = self._wrap_text(f'"{favourite_quote}"', 52)
            ty = y + 36
            for line in lines[:8]:
                draw.text((m + 28, ty), line, fill=palette['text'], font=fonts['h2'])
                ty += 26
        else:
            draw.text(
                (m + 28, y + 110),
                'No favourite quote marked yet.',
                fill=palette['muted'],
                font=fonts['body'],
            )
        return img

    def _pdf_page_timeline(self, Image, ImageDraw, fonts, palette, activity):
        img = self._pdf_blank(Image, palette)
        draw = ImageDraw.Draw(img)
        y = self._pdf_draw_section_chrome(draw, fonts, palette, 'Reading Timeline')
        m = self._PDF_MARGIN
        if not activity:
            draw.text((m, y), 'No timeline events available.', fill=palette['muted'], font=fonts['body'])
            return img
        kind_labels = {
            'finished': 'Finished',
            'started': 'Started',
            'added': 'Added',
            'wild': 'Wild Book',
        }
        for event in activity:
            ed = event['event_date']
            ed_str = ed.strftime('%Y-%m-%d') if hasattr(ed, 'strftime') else str(ed)
            book = event['book']
            cx, cy = m + 10, y + 18
            draw.ellipse((cx - 7, cy - 7, cx + 7, cy + 7), fill=palette['primary'])
            draw.line((cx, cy + 10, cx, y + 78), fill=palette['border'], width=2)
            draw.text((m + 36, y), kind_labels.get(event['kind'], event['kind']),
                      fill=palette['text'], font=fonts['body_bold'])
            draw.text(
                (m + 36, y + 32),
                self._truncate(f'{book.title} — {book.author}', 62),
                fill=palette['muted'],
                font=fonts['small'],
            )
            dw = draw.textlength(ed_str, font=fonts['small'])
            draw.text((self._PDF_W - m - dw, y), ed_str, fill=palette['muted'], font=fonts['small'])
            y += 86
            if y > self._PDF_H - 100:
                break
        return img

    def _pdf_page_recent(self, Image, ImageDraw, fonts, palette, recent):
        img = self._pdf_blank(Image, palette)
        draw = ImageDraw.Draw(img)
        y = self._pdf_draw_section_chrome(draw, fonts, palette, 'Recent Books')
        m = self._PDF_MARGIN
        if not recent:
            draw.text((m, y), 'No recent books.', fill=palette['muted'], font=fonts['body'])
            return img
        for book in recent:
            status = (
                book.reading_status
                if ReadingStatus.is_valid(book.reading_status)
                else ReadingStatus.FINISHED
            )
            added = self._fmt_datetime(getattr(book, 'date_added', None)) or '—'
            draw.text(
                (m, y),
                self._truncate(book.title, 55),
                fill=palette['text'],
                font=fonts['body_bold'],
            )
            draw.text(
                (m, y + 32),
                f'{book.author} · {ReadingStatus.label(status)} · Added {added}',
                fill=palette['muted'],
                font=fonts['small'],
            )
            y += 78
        return img

    def _pdf_page_charts(
        self,
        Image,
        ImageDraw,
        fonts,
        palette,
        *,
        rating_values,
        month_labels,
        month_values,
        year_labels,
        year_values,
    ):
        img = self._pdf_blank(Image, palette)
        draw = ImageDraw.Draw(img)
        y = self._pdf_draw_section_chrome(draw, fonts, palette, 'Charts')
        m = self._PDF_MARGIN
        chart_w = self._PDF_W - 2 * m
        y = self._pdf_draw_bar_chart(
            draw, fonts, palette,
            title='Rating Distribution',
            labels=['1', '2', '3', '4', '5'],
            values=rating_values,
            x=m, y=y, width=chart_w, height=200,
        )
        y += 36
        y = self._pdf_draw_bar_chart(
            draw, fonts, palette,
            title='Books Per Month',
            labels=month_labels,
            values=month_values,
            x=m, y=y, width=chart_w, height=200,
        )
        y += 28
        self._pdf_draw_bar_chart(
            draw, fonts, palette,
            title='Reading Throughout The Year',
            labels=year_labels,
            values=year_values,
            x=m, y=y, width=chart_w, height=200,
        )
        return img

    def _pdf_draw_bar_chart(
        self,
        draw,
        fonts,
        palette,
        *,
        title,
        labels,
        values,
        x,
        y,
        width,
        height,
    ) -> int:
        draw.text((x, y), title, fill=palette['text'], font=fonts['body_bold'])
        top = y + 36
        bottom = top + height
        draw.rounded_rectangle((x, top, x + width, bottom), radius=14, fill=palette['card'])
        if not values or not any(values):
            draw.text(
                (x + 28, top + height // 2),
                'No data for this chart.',
                fill=palette['muted'],
                font=fonts['body'],
            )
            return bottom

        max_v = max(values) or 1
        n = len(values)
        gap = 10
        inner_l = x + 28
        inner_r = x + width - 28
        plot_w = inner_r - inner_l
        bar_w = max(12, (plot_w - gap * (n - 1)) / n)
        plot_bottom = bottom - 42
        plot_top = top + 28
        plot_h = plot_bottom - plot_top
        for i, (label, value) in enumerate(zip(labels, values)):
            bh = (float(value) / max_v) * plot_h if max_v else 0
            bx = inner_l + i * (bar_w + gap)
            by = plot_bottom - bh
            draw.rounded_rectangle(
                (bx, by, bx + bar_w, plot_bottom),
                radius=6,
                fill=palette['primary'],
            )
            lw = draw.textlength(str(label), font=fonts['tiny'])
            draw.text(
                (bx + (bar_w - lw) / 2, plot_bottom + 10),
                str(label),
                fill=palette['muted'],
                font=fonts['tiny'],
            )
            if value:
                vw = draw.textlength(str(int(value)), font=fonts['tiny'])
                draw.text(
                    (bx + (bar_w - vw) / 2, by - 22),
                    str(int(value)),
                    fill=palette['text'],
                    font=fonts['tiny'],
                )
        return bottom

    # ── JSON backup ───────────────────────────────────────────

    def _export_json_backup(
        self,
        *,
        books: list[Any],
        goals: list[Any],
        username: str,
        user_created_at: datetime | None,
        avatar: str | None,
        now: datetime,
        stamp: str,
    ) -> tuple[io.BytesIO, str, str]:
        books_payload = []
        quote_total = 0
        for book in books:
            quotes = list(getattr(book, 'quotes', None) or [])
            quote_total += len(quotes)
            status = (
                book.reading_status
                if ReadingStatus.is_valid(book.reading_status)
                else ReadingStatus.FINISHED
            )
            books_payload.append({
                'id': book.id,
                'title': book.title,
                'author': book.author,
                'rating': book.rating,
                'reading_status': status,
                'is_wild': bool(book.is_wild),
                'cover_filename': book.cover_filename,
                'cover_data': book.cover_data,
                'started_reading_at': self._iso_date(book.started_reading_at),
                'finished_reading_at': self._iso_date(book.finished_reading_at),
                'date_added': self._iso_datetime(book.date_added),
                'quotes': [
                    {
                        'id': q.id,
                        'text': q.text,
                        'page_ref': q.page_ref,
                        'date_added': self._iso_datetime(q.date_added),
                        'is_favourite': bool(q.is_favourite),
                    }
                    for q in quotes
                ],
            })

        # Timeline events for reconstruction / analysis
        timeline = []
        for event in self.reading_service.recent_activity(books, limit=10_000):
            ed = event['event_date']
            timeline.append({
                'kind': event['kind'],
                'book_id': getattr(event['book'], 'id', None),
                'book_title': getattr(event['book'], 'title', None),
                'event_date': (
                    ed.isoformat() if hasattr(ed, 'isoformat') else str(ed)
                ),
            })

        payload = {
            'format': BACKUP_FORMAT,
            'version': BACKUP_VERSION,
            'exported_at': now.replace(microsecond=0).isoformat() + 'Z',
            'application': 'Kwalitec Library',
            'user': {
                'username': username,
                'created_at': self._iso_datetime(user_created_at),
                'avatar': avatar,
            },
            'settings': {
                # Theme preference is device-local (localStorage) — not stored server-side.
                'theme_note': 'Appearance theme is stored on the device only and is not part of this backup.',
                'reconstruction_notes': [
                    'Restore books with quotes nested under each book.',
                    'Ratings use 0 as the unrated sentinel; 1–5 are real scores.',
                    'reading_status values: want_to_read, reading, finished, dnf, rereading.',
                    'cover_data is a WEBP data-URI when present; cover_filename is legacy disk path.',
                    'Passwords and security answers are never included.',
                ],
            },
            'reading_goals': [
                {
                    'year': g.year,
                    'target_count': g.target_count,
                    'created_at': self._iso_datetime(getattr(g, 'created_at', None)),
                    'updated_at': self._iso_datetime(getattr(g, 'updated_at', None)),
                }
                for g in sorted(goals, key=lambda g: g.year)
            ],
            'books': books_payload,
            'timeline': timeline,
            'metadata': {
                'book_count': len(books_payload),
                'quote_count': quote_total,
                'wild_count': sum(1 for b in books if b.is_wild),
                'goal_count': len(goals),
                'favourite_quote_id': next(
                    (
                        q['id']
                        for b in books_payload
                        for q in b['quotes']
                        if q.get('is_favourite')
                    ),
                    None,
                ),
            },
        }
        raw = json.dumps(payload, ensure_ascii=False, indent=2).encode('utf-8')
        buf = io.BytesIO(raw)
        buf.seek(0)
        filename = f'kwalitec-library-backup-{stamp}.json'
        return buf, filename, 'application/json; charset=utf-8'

    # ── Helpers ───────────────────────────────────────────────

    @staticmethod
    def _fmt_date(value: date | None) -> str:
        if value is None:
            return ''
        return value.isoformat()

    @staticmethod
    def _fmt_datetime(value: datetime | None) -> str:
        if value is None:
            return ''
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return str(value)

    @staticmethod
    def _iso_date(value: date | None) -> str | None:
        if value is None:
            return None
        return value.isoformat()

    @staticmethod
    def _iso_datetime(value: datetime | None) -> str | None:
        if value is None:
            return None
        if isinstance(value, datetime):
            return value.replace(microsecond=0).isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return str(value)

    @staticmethod
    def _truncate(text: str, limit: int) -> str:
        text = text or ''
        if len(text) <= limit:
            return text
        return text[: max(0, limit - 1)] + '…'

    @staticmethod
    def _wrap_text(text: str, width: int) -> list[str]:
        words = (text or '').split()
        if not words:
            return ['']
        lines: list[str] = []
        current = words[0]
        for word in words[1:]:
            candidate = f'{current} {word}'
            if len(candidate) <= width:
                current = candidate
            else:
                lines.append(current)
                current = word
        lines.append(current)
        return lines

    @staticmethod
    def _chart_values(chart: dict[str, Any]) -> list[int]:
        raw = chart.get('data_js')
        if isinstance(raw, str):
            try:
                data = json.loads(raw)
                return [int(x) for x in data]
            except (TypeError, ValueError, json.JSONDecodeError):
                return []
        if isinstance(raw, list):
            return [int(x) for x in raw]
        return []

    def _month_chart_series(
        self, books: list[Any], now: datetime
    ) -> tuple[list[str], list[int]]:
        keys: list[str] = []
        labels: list[str] = []
        for i in range(11, -1, -1):
            y = now.year + (now.month - 1 - i) // 12
            m = (now.month - 1 - i) % 12 + 1
            keys.append(f'{y}-{m:02d}')
            labels.append(datetime(y, m, 1).strftime('%b'))
        counts: dict[str, int] = defaultdict(int)
        for b in books:
            finished = getattr(b, 'finished_reading_at', None)
            status = getattr(b, 'reading_status', None) or ReadingStatus.FINISHED
            if status == ReadingStatus.FINISHED and finished is not None:
                counts[finished.strftime('%Y-%m')] += 1
        return labels, [counts.get(k, 0) for k in keys]

    def _year_chart_series(self, books: list[Any]) -> tuple[list[str], list[int]]:
        """Finished books by calendar year (last 6 years with data, or current)."""
        counts: Counter[int] = Counter()
        for b in books:
            finished = getattr(b, 'finished_reading_at', None)
            status = getattr(b, 'reading_status', None) or ReadingStatus.FINISHED
            if status == ReadingStatus.FINISHED and finished is not None:
                counts[finished.year] += 1
        if not counts:
            year = date.today().year
            return [str(year)], [0]
        years = sorted(counts.keys())[-6:]
        return [str(y) for y in years], [counts[y] for y in years]
